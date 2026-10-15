from flask import Blueprint, render_template, request, redirect, url_for, flash, send_file
from app.utils.decorators import admin_required
from app import db
from app.models.models import Order, User, OrderStatusHistory, Notification
from sqlalchemy import func
from datetime import date, datetime
from calendar import monthrange
from io import BytesIO
import csv

admin_bp = Blueprint('admin', __name__)


def report_period():
    today = date.today()
    try:
        year, month = int(request.args.get('year', today.year)), int(request.args.get('month', today.month))
        if not (2000 <= year <= 2100 and 1 <= month <= 12):
            raise ValueError
    except (TypeError, ValueError):
        year, month = today.year, today.month
    start = date(year, month, 1)
    end = date(year + (month == 12), 1 if month == 12 else month + 1, 1)
    return year, month, start, end


def build_report(year, month, start, end):
    orders = Order.query.filter(
        Order.created_at >= datetime.combine(start, datetime.min.time()),
        Order.created_at < datetime.combine(end, datetime.min.time())
    ).order_by(Order.created_at.asc()).all()
    daily = {
        date(year, month, d).isoformat(): {'date': date(year, month, d).isoformat(),
            'orders': 0, 'sales': 0.0, 'pending': 0}
        for d in range(1, monthrange(year, month)[1] + 1)
    }
    details, paid_sales, completed, cancelled, pending = [], 0.0, 0, 0, 0
    for o in orders:
        key = o.created_at.date().isoformat()
        total = float(o.grand_total or 0)
        paid = o.payment_status == 'PAID'
        daily[key]['orders'] += 1
        if paid:
            daily[key]['sales'] += total
            paid_sales += total
        if o.status == 'COMPLETED':
            completed += 1
        if o.status in ('CANCELLED', 'REJECTED'):
            cancelled += 1
        if o.status == 'PENDING_ADMIN_REVIEW':
            daily[key]['pending'] += 1
            pending += 1
        details.append({
            'id': o.id, 'date': o.created_at.strftime('%Y-%m-%d'),
            'time': o.created_at.strftime('%I:%M %p'),
            'table': o.table.table_number if getattr(o, 'table', None) else '—',
            'status': o.status, 'payment_status': o.payment_status,
            'total': total, 'paid_sales': total if paid else 0.0
        })
    rows = list(daily.values())
    summary = {'total_orders': len(orders), 'paid_sales': round(paid_sales, 2),
        'completed_orders': completed, 'cancelled_orders': cancelled,
        'pending_orders': pending,
        'average_order_value': round(paid_sales / len(orders), 2) if orders else 0.0}
    return summary, rows, details


@admin_bp.route('/dashboard')
@admin_required
def dashboard():
    today = date.today()
    today_orders = Order.query.filter(func.date(Order.created_at) == today).count()
    sales = db.session.query(func.sum(Order.grand_total)).filter(
        func.date(Order.created_at) == today, Order.payment_status == 'PAID').scalar()
    stats = {
        'today_orders': today_orders,
        'today_sales': float(sales or 0),
        'pending_orders': Order.query.filter_by(status='PENDING_ADMIN_REVIEW').count(),
        'active_customers': User.query.filter_by(role='customer').count()
    }
    year, month, start, end = report_period()
    summary, daily_rows, _ = build_report(year, month, start, end)
    recent_orders = Order.query.order_by(Order.created_at.desc()).limit(5).all()
    return render_template('admin/dashboard.html', stats=stats, recent_orders=recent_orders,
        report_year=year, report_month=month, monthly_summary=summary, daily_rows=daily_rows,
        available_years=list(range(max(2020, today.year - 5), today.year + 1)))


@admin_bp.route('/reports/export/<file_format>')
@admin_required
def export_analytics(file_format):
    if file_format not in ('csv', 'xlsx', 'pdf'):
        flash('Unsupported report format.', 'warning')
        return redirect(url_for('admin.dashboard'))
    year, month, start, end = report_period()
    summary, daily_rows, details = build_report(year, month, start, end)
    filename = f'king_cafe_report_{year}_{month:02d}'

    if file_format == 'csv':
        # Build CSV in text memory first, then encode to bytes. This avoids
        # TextIOWrapper closing the BytesIO stream before Flask sends it.
        import io
        text_output = io.StringIO(newline='')
        writer = csv.writer(text_output)
        writer.writerow(['King Cafe Monthly Analytics', f'{year}-{month:02d}'])
        writer.writerow(['Metric', 'Value'])
        for k, v in summary.items():
            writer.writerow([k.replace('_', ' ').title(), v])
        writer.writerow([])
        writer.writerow(['Date', 'Orders', 'Paid Sales (INR)', 'Pending Review'])
        for r in daily_rows:
            writer.writerow([r['date'], r['orders'], f"{r['sales']:.2f}", r['pending']])
        writer.writerow([])
        writer.writerow(['Order ID', 'Date', 'Time', 'Table', 'Status', 'Payment Status', 'Order Total', 'Paid Sales'])
        for r in details:
            writer.writerow([r['id'], r['date'], r['time'], r['table'], r['status'],
                             r['payment_status'], f"{r['total']:.2f}", f"{r['paid_sales']:.2f}"])
        csv_bytes = ('\ufeff' + text_output.getvalue()).encode('utf-8')
        output = BytesIO(csv_bytes)
        output.seek(0)
        return send_file(output, mimetype='text/csv; charset=utf-8', as_attachment=True,
                         download_name=filename + '.csv')

    if file_format == 'xlsx':
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
            from openpyxl.utils import get_column_letter
        except ImportError:
            flash('Excel export needs openpyxl. Add openpyxl to requirements.txt and redeploy.', 'danger')
            return redirect(url_for('admin.dashboard', year=year, month=month))
        wb = Workbook()
        ws = wb.active
        ws.title = 'Monthly Summary'
        ws.append(['King Cafe Monthly Analytics', f'{year}-{month:02d}'])
        ws.append(['Metric', 'Value'])
        for k, v in summary.items():
            ws.append([k.replace('_', ' ').title(), v])
        dayws = wb.create_sheet('Daily Breakdown')
        dayws.append(['Date', 'Orders', 'Paid Sales (INR)', 'Pending Review'])
        for r in daily_rows:
            dayws.append([r['date'], r['orders'], r['sales'], r['pending']])
        detailws = wb.create_sheet('Order Details')
        detailws.append(['Order ID', 'Date', 'Time', 'Table', 'Status', 'Payment Status', 'Order Total', 'Paid Sales'])
        for r in details:
            detailws.append([r['id'], r['date'], r['time'], r['table'], r['status'],
                             r['payment_status'], r['total'], r['paid_sales']])
        for sheet in wb.worksheets:
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='50362A')
            for cells in sheet.columns:
                sheet.column_dimensions[get_column_letter(cells[0].column)].width = min(
                    max(max(len(str(c.value or '')) for c in cells) + 2, 12), 32)
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
                         as_attachment=True, download_name=filename + '.xlsx')

    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.units import mm
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
    except ImportError:
        flash('PDF export needs reportlab. Add reportlab to requirements.txt and redeploy.', 'danger')
        return redirect(url_for('admin.dashboard', year=year, month=month))
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), rightMargin=10*mm, leftMargin=10*mm)
    styles = getSampleStyleSheet()
    story = [Paragraph('King Cafe — Monthly Analytics Report', styles['Title']),
             Paragraph(f'Period: {year}-{month:02d}', styles['Normal']), Spacer(1, 6*mm),
             Paragraph('Monthly Summary', styles['Heading2'])]
    table = Table([['Metric', 'Value']] + [[k.replace('_', ' ').title(), str(v)] for k, v in summary.items()])
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#50362A')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.4,colors.grey),('PADDING',(0,0),(-1,-1),5)]))
    story += [table, Spacer(1, 5*mm), Paragraph('Daily Breakdown', styles['Heading2'])]
    daydata = [['Date', 'Orders', 'Paid Sales (INR)', 'Pending Review']]
    daydata += [[r['date'], str(r['orders']), f"{r['sales']:.2f}", str(r['pending'])] for r in daily_rows]
    table = Table(daydata, repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#50362A')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.3,colors.grey),('FONTSIZE',(0,0),(-1,-1),7)]))
    story += [table, Spacer(1, 5*mm), Paragraph('Order Details', styles['Heading2'])]
    od = [['ID', 'Date', 'Time', 'Table', 'Status', 'Payment', 'Total', 'Paid']]
    od += [[str(r['id']), r['date'], r['time'], r['table'], r['status'], r['payment_status'],
            f"{r['total']:.2f}", f"{r['paid_sales']:.2f}"] for r in details]
    if len(od) == 1:
        od.append(['—', 'No orders', '', '', '', '', '', ''])
    table = Table(od, repeatRows=1)
    table.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#50362A')),
        ('TEXTCOLOR',(0,0),(-1,0),colors.white),('GRID',(0,0),(-1,-1),.3,colors.grey),('FONTSIZE',(0,0),(-1,-1),6)]))
    story.append(table)
    doc.build(story)
    output.seek(0)
    return send_file(output, mimetype='application/pdf', as_attachment=True, download_name=filename + '.pdf')


@admin_bp.route('/orders')
@admin_required
def manage_orders():
    pending_orders = Order.query.filter_by(status='PENDING_ADMIN_REVIEW').order_by(Order.created_at.desc()).all()
    active_orders = Order.query.filter(Order.status.in_(['ACCEPTED', 'PREPARING', 'READY'])).order_by(Order.created_at.desc()).all()
    return render_template('admin/orders.html', pending_orders=pending_orders, active_orders=active_orders)


@admin_bp.route('/orders/<int:order_id>/<action>', methods=['POST'])
@admin_required
def process_order(order_id, action):
    order = Order.query.get_or_404(order_id)
    if action == 'accept':
        order.status = 'ACCEPTED'
        db.session.add(Notification(user_id=order.customer_id, role_target='customer', message=f'Your Order #{order.id} has been accepted!'))
        flash(f'Order #{order.id} accepted and sent to kitchen.', 'success')
    elif action == 'reject':
        order.status = 'REJECTED'
        reason = request.form.get('reason', 'No reason provided')
        order.notes = f"Rejected: {reason}"
        db.session.add(Notification(user_id=order.customer_id, role_target='customer', message=f'Your Order #{order.id} was rejected. Reason: {reason}'))
        flash(f'Order #{order.id} rejected.', 'warning')
    else:
        flash('Invalid order action.', 'danger')
        return redirect(url_for('admin.manage_orders'))
    db.session.add(OrderStatusHistory(order_id=order.id, status=order.status))
    db.session.commit()
    return redirect(url_for('admin.manage_orders'))


@admin_bp.route('/bill/<int:order_id>')
@admin_required
def generate_bill(order_id):
    return render_template('admin/bill.html', order=Order.query.get_or_404(order_id))


@admin_bp.route('/payment/<int:order_id>', methods=['POST'])
@admin_required
def update_payment(order_id):
    order = Order.query.get_or_404(order_id)
    payment_status = request.form.get('payment_status')
    if payment_status in ['PAID', 'PENDING', 'FAILED', 'REFUNDED']:
        order.payment_status = payment_status
        db.session.add(Notification(user_id=order.customer_id, role_target='customer',
            message=f'Payment for Order #{order.id} is now {payment_status}.'))
        if payment_status == 'PAID' and order.status in ['READY', 'SERVED']:
            order.status = 'COMPLETED'
            db.session.add(OrderStatusHistory(order_id=order.id, status='COMPLETED'))
        db.session.commit()
        flash(f'Payment status updated to {payment_status} for Order #{order.id}', 'success')
    return redirect(url_for('admin.generate_bill', order_id=order.id))


@admin_bp.route('/order-history')
@admin_required
def order_history():
    orders = Order.query.order_by(Order.created_at.desc()).all()
    return render_template('admin/order_history.html', orders=orders)
